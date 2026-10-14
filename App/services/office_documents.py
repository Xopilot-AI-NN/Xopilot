"""
Файл: /App/services/office_documents.py
Разработчик: DenBroLiik
Версия: 2.0.8
Описание: Создание настоящих DOCX, PDF, XLSX и текстовых файлов.
          Генерация документов не исполняет код или макросы.
"""
import csv
import io
from pathlib import Path
import re
from xml.sax.saxutils import escape

FORMATS = {'docx', 'pdf', 'xlsx', 'csv', 'txt', 'md', 'html', 'py', 'js', 'json', 'css', 'sql'}


def render(name, content='', *, table=None):
    name = Path(str(name).replace('\\', '/')).name
    if name in {'', '.', '..'}:
        raise ValueError('Укажите имя файла с расширением')
    extension = Path(name).suffix.lower().lstrip('.')
    if extension not in FORMATS:
        raise ValueError('Форматы: ' + ', '.join(sorted(FORMATS)))
    if not isinstance(content, str):
        raise ValueError('Содержимое документа должно быть текстом')
    if table is not None and (not isinstance(table, list) or any(not isinstance(row, list) for row in table)):
        raise ValueError('Таблица должна быть массивом строк, каждая строка — массив ячеек')
    if extension in {'xlsx', 'csv'}:
        rows = table if table is not None else list(csv.reader(io.StringIO(content)))
        if extension == 'csv':
            result = io.StringIO(); csv.writer(result).writerows(rows)
            return name, result.getvalue().encode('utf-8-sig')
        from openpyxl import Workbook
        workbook = Workbook(); sheet = workbook.active; sheet.title = 'Данные'
        for row in rows:
            sheet.append([value if isinstance(value, (str, int, float, bool)) or value is None else str(value) for value in row])
        # A cell supplied as data must not become an Excel formula/DDE link.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'
        sheet.freeze_panes = 'A2'
        if sheet.max_row and sheet.max_column:
            sheet.auto_filter.ref = sheet.dimensions
        result = io.BytesIO(); workbook.save(result); return name, result.getvalue()
    if extension == 'docx':
        from docx import Document
        document = Document()
        for line in content.splitlines():
            heading = re.match(r'^(#{1,6})\s+(.+)', line)
            if heading: document.add_heading(heading[2], level=len(heading[1]))
            else: document.add_paragraph(line)
        if table:
            grid = document.add_table(rows=0, cols=max(len(row) for row in table)); grid.style = 'Table Grid'
            for row in table:
                cells = grid.add_row().cells
                for cell, value in zip(cells, row): cell.text = str(value)
        result = io.BytesIO(); document.save(result); return name, result.getvalue()
    if extension == 'pdf':
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        font = Path(__file__).resolve().parents[1] / 'assets/fonts/GoogleSans-Regular.ttf'
        if 'XopilotSans' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('XopilotSans', str(font)))
        styles = getSampleStyleSheet()
        for style in styles.byName.values(): style.fontName = 'XopilotSans'
        story = []
        for line in content.splitlines():
            heading = re.match(r'^(#{1,6})\s+(.+)', line)
            style = styles['Heading'+str(min(3, len(heading[1])))] if heading else styles['BodyText']
            story.extend([Paragraph(escape(heading[2] if heading else line) or '&#160;', style), Spacer(1, 5)])
        if table:
            columns = max(len(row) for row in table)
            grid = Table([[Paragraph(escape(str(value)), styles['BodyText']) for value in row] +
                          [Paragraph('', styles['BodyText'])] * (columns-len(row)) for row in table],
                         colWidths=[480/columns]*columns, repeatRows=1)
            grid.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,colors.HexColor('#bfd1e5')),
                                     ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e0edff'))]))
            story.append(grid)
        result = io.BytesIO(); SimpleDocTemplate(result, title=name, leftMargin=36, rightMargin=36).build(story)
        return name, result.getvalue()
    return name, content.encode('utf-8')


def create(name, content='', *, table=None, workspace_id='all'):
    from .material_store import store_material
    from . import library
    name, data = render(name, content, table=table)
    path = store_material(name, data=data)
    if workspace_id != 'all':
        library.append_materials(workspace_id, [{'name': name, 'path': path, 'enabled': True}])
    return {'name': name, 'path': path, 'bytes': len(data), 'saved': True}
