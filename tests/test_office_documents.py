"""
Файл: /tests/test_office_documents.py
Разработчик: DenBroLiik
Версия: 2.0.8
Описание: Проверки настоящих файлов Office и поиска обычной истории.
          Исполнитель: Codex.
"""
import io
import unittest
from types import SimpleNamespace
from unittest.mock import patch
from App.services import office_documents,chat_search,workspaces


class Settings:
    def __init__(self): self.data={}
    def get_setting(self,key): return self.data.get(key)
    def set_setting(self,key,value): self.data[key]=value
    def list_chats(self): return [(1,'Задача Rust',1),(2,'Поездка',2),(-1,'Защищённый Rust',3)]
    def get_messages(self,cid):
        if cid < 0: raise AssertionError('Private history must never be accessed')
        return [SimpleNamespace(id=cid,role='user',content='Сборка Rust проекта' if cid==1 else 'Поездка на море',attachments=[])]


class OfficeDocumentsTests(unittest.TestCase):
    def test_office_files_are_readable_and_cyrillic_survives(self):
        from docx import Document
        from pypdf import PdfReader
        from openpyxl import load_workbook
        _,data=office_documents.render('Отчёт.docx','# Привет\nНастоящий документ',table=[['Название','Цена'],['Книга',42]])
        doc=Document(io.BytesIO(data));self.assertEqual(doc.paragraphs[0].text,'Привет')
        self.assertEqual(doc.tables[0].cell(1,1).text,'42')
        _,data=office_documents.render('Отчёт.pdf','# Привет\nНастоящий документ')
        self.assertIn('Привет',PdfReader(io.BytesIO(data)).pages[0].extract_text())
        _,data=office_documents.render('Данные.xlsx',table=[['Текст','Число'],['=HYPERLINK("https://example.com")',42]])
        book=load_workbook(io.BytesIO(data));self.assertEqual(book.active['A2'].data_type,'s')
        self.assertEqual(book.active['B2'].value,42)

    def test_search_and_paging_never_read_private_chats(self):
        db=Settings()
        with patch.object(chat_search,'get_db',return_value=db),patch.object(workspaces,'get_db',return_value=db):
            found=chat_search.search('Rust');self.assertEqual([row['chat_id'] for row in found['chats']],[1])
            self.assertIn('Rust',chat_search.read(1)['messages'][0]['content'])
            with self.assertRaises(ValueError):chat_search.read(-1)
            self.assertEqual(chat_search.search('Rust',after_chat_id=1)['chats'],[])
            self.assertEqual([c['chat_id'] for c in chat_search.recent(exclude_chat_id=2)['chats']],[1])
            self.assertEqual([c['chat_id'] for c in chat_search.recent()['chats']],[2,1])
